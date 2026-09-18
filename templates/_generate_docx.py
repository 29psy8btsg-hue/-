# -*- coding: utf-8 -*-
"""議事録テンプレート（Word）を生成する。"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_LINE_SPACING
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

JP = '游ゴシック'
GRAY = RGBColor(0x8A, 0x8A, 0x8A)
NAVY = RGBColor(0x1F, 0x3B, 0x63)
HEAD_FILL = 'DCE6F1'   # 見出し行の塗り
LABEL_FILL = 'F2F2F2'  # 項目名セルの塗り


def set_run(run, size=10.5, bold=False, color=None, italic=False, name=JP):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), name)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=10.5, color=None, italic=False, align=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f'{num}. {text}' if num else text)
    set_run(run, size=12, bold=True, color=NAVY)
    # 下罫線で見出しらしく
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), '1F3B63')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def hint(doc, text, size=9):
    """記入のヒント（グレー斜体・印刷前に削除して使う）"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run(run, size=size, color=GRAY, italic=True)
    return p


def body(doc, text='', size=10.5, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run(run, size=size)
    return p


def make_table(doc, rows, cols, widths, min_height=None):
    """列幅を確実に固定した表をつくる（Wordは tblLayout=fixed と tblGrid が必要）。"""
    t = doc.add_table(rows=rows, cols=cols)
    t.style = 'Table Grid'
    t.autofit = False

    tbl = t._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)

    grid = tbl.find(qn('w:tblGrid'))
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(int(round(w * 567))))  # 1cm = 567 twips
        grid.append(gc)

    for row in t.rows:
        if min_height is not None:
            row.height = Cm(min_height)
            trPr = row._tr.get_or_add_trPr()
            hr = OxmlElement('w:trHeight')
            hr.set(qn('w:val'), str(int(round(min_height * 567))))
            hr.set(qn('w:hRule'), 'atLeast')
            trPr.append(hr)
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
    return t


def header_row(table, r, labels):
    for i, label in enumerate(labels):
        c = table.cell(r, i)
        shade(c, HEAD_FILL)
        cell_text(c, label, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)


def build(path):
    doc = Document()

    # --- 既定フォント（日本語対応） ---
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.font.name = JP
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), JP)

    # --- ページ設定（A4 縦・余白狭め） ---
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)
    usable = 21.0 - 1.9 * 2  # 17.2cm

    # --- フッター（ページ番号は手動で十分なので文書名のみ） ---
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer_p.add_run('議事録'), size=9, color=GRAY)

    # --- タイトル ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_run(title.add_run('議 事 録'), size=20, bold=True, color=NAVY)

    # --- 基本情報 ---
    w = [3.2, 5.4, 3.2, 5.4]
    t = make_table(doc, 6, 4, w, min_height=0.7)
    rows = [
        ('会議名', '', '文書番号', ''),
        ('開催日時', '', '開催場所', ''),
        ('主催・議長', '', '書記・作成者', ''),
    ]
    for r, (l1, v1, l2, v2) in enumerate(rows):
        shade(t.cell(r, 0), LABEL_FILL)
        shade(t.cell(r, 2), LABEL_FILL)
        cell_text(t.cell(r, 0), l1, bold=True, size=10)
        cell_text(t.cell(r, 1), v1)
        cell_text(t.cell(r, 2), l2, bold=True, size=10)
        cell_text(t.cell(r, 3), v2)

    wide = [('出席者', ''), ('欠席者', ''), ('配布先', '')]
    for i, (label, val) in enumerate(wide):
        r = 3 + i
        shade(t.cell(r, 0), LABEL_FILL)
        cell_text(t.cell(r, 0), label, bold=True, size=10)
        merged = t.cell(r, 1).merge(t.cell(r, 3))
        cell_text(merged, val)

    hint(doc, '※ 開催日時は「2026/09/18（金）10:00〜11:00」のように曜日まで入れると後から探しやすくなります。')

    # --- 1. 目的・ゴール ---
    heading(doc, '会議の目的・ゴール', 1)
    hint(doc, '（この会議で「何が決まれば成功か」を1〜2行で。例：新プランの価格帯を1案に絞る）')
    t = make_table(doc, 1, 1, [usable], min_height=1.4)
    cell_text(t.cell(0, 0), '')

    # --- 2. アジェンダ ---
    heading(doc, 'アジェンダ（議題）', 2)
    t = make_table(doc, 4, 4, [1.4, 9.4, 3.2, 3.2], min_height=0.75)
    header_row(t, 0, ['No', '議題', '担当', '予定時間'])
    for r in range(1, 4):
        cell_text(t.cell(r, 0), str(r), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in range(1, 4):
            cell_text(t.cell(r, c), '')

    # --- 3. 決定事項 ---
    heading(doc, '決定事項', 3)
    hint(doc, '（決まったことだけを書く。決まらなかったことは「6. 保留・継続検討事項」へ）')
    t = make_table(doc, 4, 4, [1.4, 7.4, 5.2, 3.2], min_height=0.95)
    header_row(t, 0, ['No', '決定内容', '決定理由・背景', '決定者'])
    for r in range(1, 4):
        cell_text(t.cell(r, 0), f'D-{r}', align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in range(1, 4):
            cell_text(t.cell(r, c), '')

    # --- 4. アクションアイテム ---
    heading(doc, 'アクションアイテム（ToDo）', 4)
    hint(doc, '（「誰が・何を・いつまでに」の3点を必ず埋める。担当が空欄のタスクは実行されません）')
    t = make_table(doc, 5, 5, [1.4, 7.4, 2.7, 2.9, 2.8], min_height=0.95)
    header_row(t, 0, ['No', '対応内容', '担当者', '期限', '状況'])
    for r in range(1, 5):
        cell_text(t.cell(r, 0), f'A-{r}', align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(r, 1), '')
        cell_text(t.cell(r, 2), '')
        cell_text(t.cell(r, 3), '')
        cell_text(t.cell(r, 4), '未着手', align=WD_ALIGN_PARAGRAPH.CENTER)

    # --- 5. 議事内容 ---
    heading(doc, '議事内容（詳細）', 5)
    hint(doc, '（アジェンダごとに「議論の要点 → 結論」の順で。発言の逐語録ではなく論点を残す）')
    for i in (1, 2, 3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(f'議題{i}：'), size=11, bold=True)
        set_run(p.add_run('（アジェンダNo.%d のタイトル）' % i), size=10, color=GRAY, italic=True)
        t = make_table(doc, 2, 2, [2.6, usable - 2.6], min_height=1.1)
        shade(t.cell(0, 0), LABEL_FILL)
        shade(t.cell(1, 0), LABEL_FILL)
        cell_text(t.cell(0, 0), '議論の要点', bold=True, size=10)
        cell_text(t.cell(1, 0), '結論', bold=True, size=10)
        cell_text(t.cell(0, 1), '')
        cell_text(t.cell(1, 1), '')

    # --- 6. 保留 ---
    heading(doc, '保留・継続検討事項', 6)
    t = make_table(doc, 3, 4, [1.4, 8.4, 3.7, 3.7], min_height=0.9)
    header_row(t, 0, ['No', '内容', '保留の理由', '再検討の時期'])
    for r in range(1, 3):
        cell_text(t.cell(r, 0), f'P-{r}', align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in range(1, 4):
            cell_text(t.cell(r, c), '')

    # --- 7. 次回会議 ---
    heading(doc, '次回会議の予定', 7)
    t = make_table(doc, 2, 4, [3.2, 5.4, 3.2, 5.4], min_height=0.75)
    for r, (l1, l2) in enumerate([('日時', '場所'), ('主な議題', '事前準備')]):
        shade(t.cell(r, 0), LABEL_FILL)
        shade(t.cell(r, 2), LABEL_FILL)
        cell_text(t.cell(r, 0), l1, bold=True, size=10)
        cell_text(t.cell(r, 1), '')
        cell_text(t.cell(r, 2), l2, bold=True, size=10)
        cell_text(t.cell(r, 3), '')

    # --- 8. 添付資料 ---
    heading(doc, '添付資料・参考リンク', 8)
    body(doc, '・')
    body(doc, '・')

    # --- 使い方（印刷前に削除） ---
    doc.add_page_break()
    heading(doc, 'このテンプレートの使い方（配布前にこのページごと削除してください）')
    notes = [
        '1) このファイルをコピーして1会議＝1ファイルで使います。ファイル名は「20260918_定例会議_議事録.docx」のように日付を先頭に置くと時系列で並びます。',
        '2) 会議中は「3. 決定事項」と「4. アクションアイテム」だけをその場で埋めます。この2つが埋まっていれば議事録としては成立します。',
        '3) 「5. 議事内容」は会議後に清書します。発言の書き起こしではなく、論点・出た案・選んだ理由を残してください。',
        '4) グレーの斜体はすべて記入のヒントです。書き終えたら削除してください。',
        '5) 表の行が足りないときは、最終行の右端セルにカーソルを置いて Tab キーを押すと行が増えます。',
        '6) 配布は編集不可のPDFで（ファイル → 名前を付けて保存 → PDF）。原本のWordは自分の手元に残します。',
        '7) アクションアイテムの進捗を複数の会議にまたがって追いかけたい場合は、同梱のExcel版テンプレートの「アクション管理」シートを併用してください。',
    ]
    for n in notes:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Cm(0.5)
        set_run(p.add_run(n), size=10)

    doc.save(path)
    print('saved:', path)


if __name__ == '__main__':
    build(sys.argv[1])
