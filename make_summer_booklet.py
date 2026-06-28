#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""夏休みのしおり（放課後等デイサービス 夢門塾ゆうゆう日吉）Word生成"""
import os
from docx import Document
from docx.shared import Pt, Mm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ASSET = "/home/user/-/assets"
LOGO  = os.path.join(ASSET, "logo.png")
ICO   = os.path.join(ASSET, "ico")
OUT   = "/home/user/-/夏休みのしおり_2026.docx"

FONT = "BIZ UDPゴシック"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
ORANGE = RGBColor(0xE0, 0x78, 0x20)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
PINK = RGBColor(0xD6, 0x4D, 0x7A)
GRAY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# ── 既定フォント ──
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

# ── ページ設定（A4縦・余白）──
sec = doc.sections[0]
sec.page_width  = Mm(210)
sec.page_height = Mm(297)
sec.top_margin    = Mm(15)
sec.bottom_margin = Mm(15)
sec.left_margin   = Mm(18)
sec.right_margin  = Mm(18)

# ── ヘルパー ──
def set_run(r, size=11, bold=False, color=None, font=FONT):
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), font)

def para(text="", size=11, bold=False, color=None, align=None,
         space_after=6, space_before=0, line=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if line: pf.line_spacing = line
    if text:
        r = p.add_run(text)
        set_run(r, size, bold, color)
    return p

def shade_cell(cell, hex6):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex6)
    tcPr.append(sh)

def set_cell_text(cell, text, size=11, bold=False, color=None,
                  align=WD_ALIGN_PARAGRAPH.LEFT, valign=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = valign
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text)
    set_run(r, size, bold, color)
    return cell

def set_cell_borders(table, color="BBBBBB", sz="6"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)

def section_heading(text, color=NAVY, icon=None):
    """見出しバー（左に色帯）"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    # 左帯（■）+ 見出し
    r0 = p.add_run("■ ")
    set_run(r0, 14, True, color)
    r = p.add_run(text)
    set_run(r, 14, True, color)
    # 下線（段落罫線）
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), f"{color}" if isinstance(color,str) else "%02X%02X%02X"%(color[0],color[1],color[2]))
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p

# ══════════════════════════════════════════════════════════════════
#  1) 表紙
# ══════════════════════════════════════════════════════════════════
# 上部スペース
para("", space_after=10)
para("", space_after=10)

# ロゴ
pic_p = doc.add_paragraph()
pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = pic_p.add_run()
run.add_picture(LOGO, width=Mm(90))

para("", space_after=6)

# タイトル
para("夏休みのしおり", size=40, bold=True, color=ORANGE,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para("〜 2026年（令和8年）夏 〜", size=18, bold=True, color=NAVY,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)

# 装飾アイコン（ひまわり的に balloon / star など）
deco = doc.add_paragraph()
deco.alignment = WD_ALIGN_PARAGRAPH.CENTER
for ico in ["balloon.png", "burger.png", "bowling.png", "kouen.png", "star.png"]:
    path = os.path.join(ICO, ico)
    if os.path.exists(path):
        r = deco.add_run()
        r.add_picture(path, width=Mm(18))
        deco.add_run("  ")

para("", space_after=24)

# 施設名ボックス
tbl = doc.add_table(rows=1, cols=1)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.columns[0].width = Mm(120)
c = tbl.cell(0, 0)
c.width = Mm(120)
shade_cell(c, "FFF3E0")
set_cell_text(c, "放課後等デイサービス\n夢門塾ゆうゆう日吉",
              size=16, bold=True, color=NAVY,
              align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_borders(tbl, color="E07820", sz="12")

para("", space_after=10)
para("夏休み期間：2026年7月21日（火）〜 8月31日（月）",
     size=12, bold=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

# 改ページ
doc.add_page_break()

# ══════════════════════════════════════════════════════════════════
#  2) ごあいさつ
# ══════════════════════════════════════════════════════════════════
section_heading("ごあいさつ", color=NAVY)
greeting = (
    "保護者の皆さま、いつも夢門塾ゆうゆう日吉をご利用いただき、誠にありがとうございます。\n"
    "いよいよ子どもたちが楽しみにしている夏休みがやってきます。長いお休みの間、"
    "お子さま一人ひとりが安全に、そして元気いっぱい過ごせるよう、職員一同心を込めて"
    "支援してまいります。\n"
    "今年の夏も、季節の行事やお出かけ、クッキングなど、わくわくする活動をたくさん"
    "用意しています。このしおりには、夏休み中の過ごし方や持ち物、お約束などをまとめました。"
    "ご家庭でもお子さまと一緒にご確認いただき、楽しい夏休みにしましょう。\n"
    "暑い日が続きますので、体調管理にはくれぐれもお気をつけください。"
)
for line in greeting.split("\n"):
    para(line, size=11, line=1.4, space_after=4)

# ══════════════════════════════════════════════════════════════════
#  3) 一日の流れ
# ══════════════════════════════════════════════════════════════════
section_heading("夏休みの一日の流れ", color=GREEN)
para("※開所時間：平日 8:30 〜 17:30（送迎時間は別途ご相談ください）",
     size=10, color=GRAY, space_after=6)

schedule = [
    ("8:30",  "順次登所・健康チェック（検温・手洗い・うがい）"),
    ("9:30",  "朝の会・ラジオ体操・今日の予定確認"),
    ("10:00", "学習タイム（宿題・自主学習）"),
    ("11:00", "午前の活動（運動・製作・個別療育）"),
    ("12:00", "昼食・歯みがき"),
    ("13:00", "休憩・自由遊び"),
    ("14:00", "午後の活動（行事・お出かけ・グループ活動）"),
    ("15:30", "おやつ・片付け・掃除"),
    ("16:00", "帰りの会・順次降所・送迎"),
    ("17:30", "閉所"),
]
t = doc.add_table(rows=len(schedule)+1, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.columns[0].width = Mm(28)
t.columns[1].width = Mm(146)
set_cell_borders(t, color="A5D6A7", sz="6")
# ヘッダー
set_cell_text(t.cell(0,0), "時間", bold=True, color=RGBColor(0xFF,0xFF,0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell_text(t.cell(0,1), "活動内容", bold=True, color=RGBColor(0xFF,0xFF,0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
shade_cell(t.cell(0,0), "2E7D32"); shade_cell(t.cell(0,1), "2E7D32")
for i,(time,act) in enumerate(schedule, start=1):
    t.cell(i,0).width = Mm(28); t.cell(i,1).width = Mm(146)
    set_cell_text(t.cell(i,0), time, bold=True, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t.cell(i,1), act)
    if i % 2 == 0:
        shade_cell(t.cell(i,0), "F1F8E9"); shade_cell(t.cell(i,1), "F1F8E9")

# ══════════════════════════════════════════════════════════════════
#  4) 持ち物リスト
# ══════════════════════════════════════════════════════════════════
doc.add_page_break()
section_heading("毎日の持ち物", color=ORANGE)
para("登所する日は、下記の持ち物をご準備ください。すべての持ち物に記名をお願いします。",
     size=11, space_after=6)

items = [
    "お弁当・水筒（多めの水分を）",
    "汗ふきタオル・ハンカチ・ティッシュ",
    "着替え一式（下着・靴下含む）",
    "帽子（外出・外遊び用）",
    "夏休みの宿題・学習道具",
    "連絡帳・お薬（必要な方）",
    "上ばき（室内用）",
    "ビニール袋（汚れ物入れ）",
]
ti = doc.add_table(rows=(len(items)+1)//2, cols=2)
ti.alignment = WD_TABLE_ALIGNMENT.CENTER
set_cell_borders(ti, color="FFCC80", sz="6")
for idx, item in enumerate(items):
    r, cidx = idx//2, idx%2
    cell = ti.cell(r, cidx)
    cell.width = Mm(87)
    set_cell_text(cell, "☑ " + item, size=11)
    shade_cell(cell, "FFF8E1")

para("", space_after=4)
para("★ プールや水遊びの日は、別途「水着・タオル・水泳帽」をお持ちください（前日にお知らせします）。",
     size=10, bold=True, color=PINK, space_after=4)

# ══════════════════════════════════════════════════════════════════
#  5) 夏休みのお約束
# ══════════════════════════════════════════════════════════════════
section_heading("夏休みのお約束", color=PINK)
rules = [
    "早ね・早おき・朝ごはんで、生活リズムをととのえよう。",
    "こまめに水分をとって、熱中症をふせごう。",
    "外に出るときは帽子をかぶろう。",
    "お友だちや先生の話を、最後まで聞こう。",
    "使ったものは、自分でかたづけよう。",
    "こまったときは、すぐに先生に伝えよう。",
    "交通ルールを守って、安全に過ごそう。",
]
for i, rule in enumerate(rules, start=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"{i}. ")
    set_run(r1, 12, True, PINK)
    r2 = p.add_run(rule)
    set_run(r2, 12, False)

# ══════════════════════════════════════════════════════════════════
#  6) 夏休みの主な行事予定
# ══════════════════════════════════════════════════════════════════
doc.add_page_break()
section_heading("夏休みの主な行事予定", color=NAVY)
para("※天候や状況により、内容が変更になる場合があります。詳しくは毎月の予定表でお知らせします。",
     size=10, color=GRAY, space_after=6)

events = [
    ("7月", "夏まつりごっこ", "ヨーヨーつり・的あてなど、みんなでお店やさんごっこ"),
    ("7月", "水遊び・プール活動", "暑さに負けず、お水でリフレッシュ"),
    ("8月", "クッキング（流しそうめん・かき氷）", "夏ならではの食を楽しもう"),
    ("8月", "工作week（うちわ・風鈴づくり）", "夏の作品づくりに挑戦"),
    ("8月", "お出かけ（公園・アスレチック）", "体をいっぱい動かそう"),
    ("8月", "映画・お楽しみ会", "DVD鑑賞やゲーム大会でのんびり"),
]
te = doc.add_table(rows=len(events)+1, cols=3)
te.alignment = WD_TABLE_ALIGNMENT.CENTER
te.columns[0].width = Mm(18)
te.columns[1].width = Mm(66)
te.columns[2].width = Mm(90)
set_cell_borders(te, color="9FC5E8", sz="6")
for j,h in enumerate(["時期","行事","内容"]):
    set_cell_text(te.cell(0,j), h, bold=True, color=RGBColor(0xFF,0xFF,0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(te.cell(0,j), "1F4E79")
for i,(month,name,desc) in enumerate(events, start=1):
    te.cell(i,0).width=Mm(18); te.cell(i,1).width=Mm(66); te.cell(i,2).width=Mm(90)
    set_cell_text(te.cell(i,0), month, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(te.cell(i,1), name, bold=True)
    set_cell_text(te.cell(i,2), desc, size=10)
    if i % 2 == 0:
        for j in range(3): shade_cell(te.cell(i,j), "EAF1F8")

# ══════════════════════════════════════════════════════════════════
#  7) お願い・緊急連絡
# ══════════════════════════════════════════════════════════════════
section_heading("ご家庭へのお願い", color=GREEN)
asks = [
    "欠席・遅刻・送迎時間の変更は、当日 9:00 までにご連絡ください。",
    "発熱（37.5℃以上）や体調不良がある場合は、利用をお控えください。",
    "お薬を預ける場合は、連絡帳に服用方法をご記入のうえお渡しください。",
    "持ち物には必ずお名前のご記入をお願いします。",
    "高価なもの・ゲーム機などはお持たせにならないようご協力ください。",
]
for a in asks:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run("・")
    set_run(r1, 11, True, GREEN)
    r2 = p.add_run(a)
    set_run(r2, 11)

para("", space_after=8)

# 連絡先ボックス
tb = doc.add_table(rows=1, cols=1)
tb.alignment = WD_TABLE_ALIGNMENT.CENTER
tb.columns[0].width = Mm(174)
cc = tb.cell(0,0); cc.width = Mm(174)
shade_cell(cc, "E8F5E9")
set_cell_borders(tb, color="2E7D32", sz="12")
cc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
p = cc.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("【緊急連絡先】 放課後等デイサービス 夢門塾ゆうゆう日吉")
set_run(r, 12, True, NAVY)
p2 = cc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("TEL：000-0000-0000　／　受付時間 平日 8:30〜17:30")
set_run(r2, 12, True, GRAY)

para("", space_after=10)
para("楽しい夏休みになりますように！　職員一同", size=12, bold=True,
     color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print("保存:", OUT, os.path.getsize(OUT), "bytes")
